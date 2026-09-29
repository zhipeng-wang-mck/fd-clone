"""Convert an Excel workbook into markdown tables, one per sheet.

Workbook chrome such as a Metadata Field/Value sheet is skipped so it
does not leak into the knowledge draft.
"""
from __future__ import annotations

from pathlib import Path

_SKIP_SHEET_NAMES = {"metadata", "meta", "properties"}


def _cell(value) -> str:
    if value is None:
        return ""
    return str(value).replace("|", "\\|").replace("\n", " ")


def _is_metadata_sheet(title: str, rows: list) -> bool:
    if title.casefold() in _SKIP_SHEET_NAMES:
        return True
    if not rows:
        return False
    header = [str(c or "").strip().casefold() for c in rows[0]]
    return header[:2] == ["field", "value"] and len(rows) <= 8


def xlsx_to_markdown(path: Path) -> tuple[str, list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError("openpyxl is required for XLSX conversion. Run setup.py.") from exc

    wb = load_workbook(path, data_only=True, read_only=True)
    parts: list[str] = []
    skipped: list[str] = []
    for sheet in wb.worksheets:
        title = sheet.title or "Sheet"
        rows = list(sheet.iter_rows(values_only=True))
        if _is_metadata_sheet(title, rows):
            skipped.append(title)
            continue
        parts.append(f"## {title}")
        if not rows:
            parts.append("(empty sheet)")
            parts.append("")
            continue

        header = rows[0]
        parts.append("| " + " | ".join(_cell(c) for c in header) + " |")
        parts.append("| " + " | ".join("---" for _ in header) + " |")
        for row in rows[1:]:
            parts.append("| " + " | ".join(_cell(c) for c in row) + " |")
        parts.append("")

    sheet_count = len(wb.sheetnames)
    wb.close()
    notes = [f"Converted from XLSX with openpyxl, {sheet_count} sheet(s)"]
    if skipped:
        notes.append("Skipped metadata sheet(s): " + ", ".join(skipped))
    if not parts:
        return "", notes + ["No playbook sheets remained after filtering"]
    return "\n".join(parts), notes
